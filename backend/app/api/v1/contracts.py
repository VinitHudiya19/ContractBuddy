"""
Contract endpoints.

Upload a contract, get structured analysis back (parties, obligations, missing
clauses, risk/health scores), list them, and re-run the analysis on demand.
Every query is scoped to the calling user. A contract id from another account
returns 404, not someone else's data.
"""
from __future__ import annotations

from pathlib import Path
from uuid import UUID, uuid4

from fastapi import APIRouter, Depends, File, Form, UploadFile, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.exceptions import (
    NotFoundError,
    UnsupportedFileTypeError,
    ValidationError,
)
from app.core.logging import get_logger
from app.db.session import get_db
from app.dependencies.auth import get_current_user
from app.dependencies.rate_limit import ai_rate_limit
from app.models.contract import Contract, ContractVersion
from app.models.user import User
from app.repositories.contract_repo import ContractRepository
from app.schemas.contract import ContractResponse
from app.services.contract_analysis import analyze_contract
from app.services.uploads import save_upload

logger = get_logger(__name__)

router = APIRouter(prefix="/api/contracts", tags=["contracts"])

_SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


def _contract_dir() -> Path:
    path = settings.upload_path / "contracts"
    path.mkdir(parents=True, exist_ok=True)
    return path


def extract_text(path: Path) -> str:
    """Best-effort plain text for analysis. Returns '' if nothing can be read."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            import fitz

            with fitz.open(path) as doc:
                return "\n".join(page.get_text() for page in doc)
        if suffix == ".docx":
            import docx

            document = docx.Document(str(path))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if cells:
                        parts.append(cells)
            return "\n".join(parts)
        if suffix in (".txt", ".md"):
            return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as exc:
        logger.warning(
            "contract text extraction failed",
            extra={"file": path.name, "error": f"{type(exc).__name__}: {exc}"},
        )
    return ""


@router.post(
    "",
    response_model=ContractResponse,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(ai_rate_limit)],
)
@router.post(
    "/",
    response_model=ContractResponse,
    status_code=status.HTTP_201_CREATED,
    include_in_schema=False,
    dependencies=[Depends(ai_rate_limit)],
)
async def create_contract(
    title: str = Form(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> ContractResponse:
    """Store a contract file and run structured analysis over its text."""
    if not title.strip():
        raise ValidationError("Contract title is required.")

    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in _SUPPORTED:
        raise UnsupportedFileTypeError(
            f"Supported contract formats: {', '.join(sorted(_SUPPORTED))}."
        )

    path = _contract_dir() / f"{uuid4()}{suffix}"
    await save_upload(file, path)

    text = extract_text(path)
    analysis = await analyze_contract(title, text)

    repo = ContractRepository(db)
    contract = await repo.create_contract(
        title=title.strip(),
        filename=file.filename or path.name,
        organization_id=user.id,
        **analysis,
    )
    await repo.add_version(contract, str(path))
    await db.commit()
    await db.refresh(contract)

    logger.info(
        "contract analysed",
        extra={
            "contract_id": str(contract.id),
            "source": analysis.get("analysis_source"),
            "chars": len(text),
        },
    )
    return ContractResponse.model_validate(contract)


@router.get("", response_model=list[ContractResponse])
@router.get("/", response_model=list[ContractResponse], include_in_schema=False)
async def list_contracts(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> list[ContractResponse]:
    """Every contract belonging to the current user, newest first."""
    result = await db.execute(
        select(Contract)
        .where(Contract.organization_id == user.id)
        .order_by(Contract.created_at.desc())
    )
    return [ContractResponse.model_validate(c) for c in result.scalars().all()]


@router.get("/{contract_id}", response_model=ContractResponse)
async def get_contract(
    contract_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> ContractResponse:
    contract = await _get_owned(db, contract_id, user)
    return ContractResponse.model_validate(contract)


@router.delete("/{contract_id}", status_code=204, response_model=None)
async def delete_contract(
    contract_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> None:
    contract = await _get_owned(db, contract_id, user)

    # Remove the stored files too, so deletion is not just a database row.
    versions = (
        (
            await db.execute(
                select(ContractVersion).where(ContractVersion.contract_id == contract.id)
            )
        )
        .scalars()
        .all()
    )
    for version in versions:
        try:
            Path(version.file_path).unlink(missing_ok=True)
        except OSError as exc:
            logger.warning(
                "could not delete contract file",
                extra={"path": version.file_path, "error": str(exc)},
            )

    await db.delete(contract)
    await db.commit()


@router.post(
    "/{contract_id}/analyze",
    response_model=ContractResponse,
    dependencies=[Depends(ai_rate_limit)],
)
async def reanalyze_contract(
    contract_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> ContractResponse:
    """
    Re-run analysis against the stored file. Useful after adding an API key,
    since the first pass may have used the rule-based fallback.
    """
    contract = await _get_owned(db, contract_id, user)

    version = (
        (
            await db.execute(
                select(ContractVersion)
                .where(ContractVersion.contract_id == contract.id)
                .order_by(ContractVersion.created_at.desc())
                .limit(1)
            )
        )
        .scalars()
        .first()
    )

    text = ""
    if version and Path(version.file_path).exists():
        text = extract_text(Path(version.file_path))
    else:
        logger.warning(
            "stored contract file missing, analysing from title only",
            extra={"contract_id": str(contract.id)},
        )

    analysis = await analyze_contract(contract.title, text)
    repo = ContractRepository(db)
    updated = await repo.update_analysis(contract, analysis)
    return ContractResponse.model_validate(updated)


async def _get_owned(db: AsyncSession, contract_id: UUID, user: User) -> Contract:
    result = await db.execute(
        select(Contract).where(
            Contract.id == contract_id, Contract.organization_id == user.id
        )
    )
    contract = result.scalars().first()
    if contract is None:
        raise NotFoundError("Contract not found.")
    return contract
